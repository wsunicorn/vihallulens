"""Build docs/BaoCao_KLTN.docx from docs/BaoCao_KLTN.md, following the Khoa template.

docs/BaoCao_KLTN.md itself is generated: docs/tools/resolve.py assembles it from the
chapter drafts in docs/report/, which carry labels instead of numbers. Edit the drafts,
run resolve.py, then this script — see docs/report/README.md.

The numbers come from chapter 2 of the Khoa template, "docs/Mau bao cao KLTN_DS.docx"
(its text is also in docs/Mau_bao_cao_KLTN_DS.md):

* A4, margins top 3 cm / bottom 3 cm / left 3.5 cm / right 2 cm, Times New Roman.
* Content: 13 pt, before 10 pt, after 0, 1.5 lines, justified, no first-line indent.
* Heading 1: 14 pt bold upper-case, 24/24 pt, 1.15 lines, left. Heading 2: 13 bold,
  6/12 pt. Heading 3: 13 bold italic. Heading 4: 13 italic.
* Caption: 13 pt, 6/12 pt, 1.15 lines, centred; table captions above, figure captions
  below, equation numbers in parentheses right-aligned.
* Lists of figures/tables: 13 pt, 0/0 pt, 1.15 lines. Bullets: 13 pt, 0/0, 1.5 lines.
* Page numbers bottom right, starting at 1 on the first page of Chapter 1; footer carries
  the students, header the thesis ("Header thể hiện thông tin khóa luận tốt nghiệp,
  footer thể hiện thông tin sinh viên thực hiện"). TOC down to four-digit sub-sections.
  The footer runs from the ABSTRACT page on, so the front matter prints roman numbers
  and the roman page numbers the TOC gives for them can be looked up; the header starts with
  Chapter 1. The covers stay bare.
* Printed on one side of the sheet, as the template's text says ("trình bày trên một mặt
  giấy"): the 3,5 cm binding margin on the left of every page, the page number bottom
  right, the header flush left. The two-sided layout tried on 28/09 was dropped on 09/10;
  only the landscape page keeps its margins from it (see LAND_MARGIN).

Everything chapter 2 leaves unsaid is copied from how the template file itself looks:
the style names (Content, Bullet, CoverB, Heading 1N), chapter and section numbers drawn
by a Word multilevel list ("CHƯƠNG 1" + tab, "1.1" + tab, hanging indent), italic
captions, centred unnumbered titles for the lists and the back matter, the covers (20 pt
title, art page border on the outer cover, a plain frame on the English one) and the
13 pt footer.

Pipeline: the Markdown is pre-processed into pandoc-friendly form, pandoc produces a raw
.docx (inline $...$ math becomes native Word equations there), python-docx applies every
rule above, and Word itself then builds each numbered display equation from its LaTeX
source, updates the fields — table of contents, lists of figures and tables, caption and
equation numbers, page numbers — and exports a PDF.

Display equations: each $$ ... $$ block becomes a borderless three-column table — empty
cell, the equation, and "(chapter.SEQ)" right-aligned. Word's own LaTeX converter
(Insert > Equation > LaTeX, driven through COM) turns the source into OMML, so the result
is exactly what typing the same LaTeX into Word gives. A "\\" inside a block starts a new
line of the same equation. Word must have its equation input set to LaTeX; the build
stops if any backslash survives the conversion.

Run:  python docs/tools/build_docx.py
"""

from __future__ import annotations

import copy
import math
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image

DOCS = Path(__file__).resolve().parents[1]
SRC_MD = DOCS / "BaoCao_KLTN.md"
OUT_DOCX = DOCS / "BaoCao_KLTN.docx"
LOGO = DOCS / "media" / "logo_iuh.png"

FONT = "Times New Roman"
# list marker: a Times New Roman hyphen, not the Symbol-font round dot pandoc emits.
# The Khoa template carries this exact definition — "-" (U+002D) in Times New Roman.
BULLET = "-"
PAGE_W_CM, PAGE_H_CM = 21.0, 29.7
MARGIN = dict(top=3.0, bottom=3.0, left=3.5, right=2.0)
TEXT_W_CM = PAGE_W_CM - MARGIN["left"] - MARGIN["right"]      # 15.5 cm
TEXT_H_CM = PAGE_H_CM - MARGIN["top"] - MARGIN["bottom"]      # 23.7 cm
CAPTION_CM = 1.3      # a one-line caption above a table, 6/12 pt around it
TABLE_PT = 11          # "các đề mục khác": Times New Roman 10–13, thống nhất

# ---------------------------------------------------------------- cover data ---
TITLE_VI = ("NGHIÊN CỨU VÀ XÂY DỰNG HỆ THỐNG PHÁT HIỆN ẢO GIÁC CHO MÔ HÌNH NGÔN NGỮ LỚN "
            "TĂNG CƯỜNG TRUY XUẤT TIẾNG VIỆT DỰA TRÊN TÍN HIỆU CHÚ Ý NỘI TẠI")
TITLE_EN = ("RESEARCH AND DEVELOPMENT OF A HALLUCINATION DETECTION SYSTEM FOR VIETNAMESE "
            "RETRIEVAL-AUGMENTED LARGE LANGUAGE MODELS BASED ON INTERNAL ATTENTION SIGNALS")
STUDENTS_VI = ["NGUYỄN NGỌC LÂN — MSSV 22635801", "NGUYỄN TẤN MINH — MSSV 22643511"]
STUDENTS_EN = ["NGUYEN NGOC LAN — 22635801", "NGUYEN TAN MINH — 22643511"]
# the two students in the footer, in plain 13 pt — it has to fit on one line of the
# 15,5 cm text width next to the right-aligned page number
FOOTER_TEXT = "Nguyễn Ngọc Lân – Nguyễn Tấn Minh"
# the thesis in the header, 13 pt like the template's own; at 427 pt it still fits the
# 439 pt text width on one line
HEADER_TEXT = "Khóa luận tốt nghiệp · Phát hiện ảo giác cho LLM tăng cường truy xuất tiếng Việt"
TOC1_HANG_CM = 2.9633          # 1.680 twips: the tab Word writes after "CHƯƠNG n"

FRONT_TITLES = ("ABSTRACT", "TÓM TẮT", "LỜI CẢM ƠN", "NHẬN XÉT CỦA GIÁO VIÊN HƯỚNG DẪN",
                "NHẬN XÉT CỦA GIÁO VIÊN PHẢN BIỆN 1", "NHẬN XÉT CỦA GIÁO VIÊN PHẢN BIỆN 2",
                "MỤC LỤC")
# page titles that look like a chapter but carry no number — "Heading 1N" in the
# template, centred, still listed in the table of contents
UNNUMBERED_TITLES = ("DANH MỤC", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC", "NHẬT KÍ LÀM VIỆC")
PAGEBREAK_DIV = '<div style="page-break-after: always"></div>'

# charts made with matplotlib read better a little narrower than diagrams/screenshots
NARROW_IMAGES = ("bieudo_",)
# dense diagrams go on a landscape page — "hình vẽ trình bày theo chiều ngang khổ giấy
# thì đầu bảng là lề trái của trang"
LANDSCAPE_IMAGES = ("hinh_trich_huanluyen_phucvu",)
# A landscape page is turned onto the sheet, so its long top and bottom edges lie along
# the spine and the outer edge. The binding margin on both, with header and footer kept
# clear of it, dates from the two-sided layout, when either edge could meet the spine;
# the book went back to one side on 09/10 and kept these margins (and with them the
# 19,2 cm width of Hình 5.1) rather than widen the figure again.
LAND_MARGIN = dict(top=3.5, bottom=3.5, left=3.0, right=3.0)
LAND_HF_CM = 2.2                 # header/footer distance from the edge on that page
LAND_TEXT_W_CM = PAGE_H_CM - LAND_MARGIN["left"] - LAND_MARGIN["right"]   # 23.7 cm
LAND_TEXT_H_CM = PAGE_W_CM - LAND_MARGIN["top"] - LAND_MARGIN["bottom"]   # 14.0 cm


# =============================================================================
# 1. Pre-processing
# =============================================================================
def shrink_image(src: Path, dst_dir: Path, max_px: int = 1800) -> Path:
    """Copy an image, downsampled to at most max_px wide (~295 dpi at 15.5 cm)."""
    dst = dst_dir / src.name
    im = Image.open(src)
    if im.width > max_px:
        im = im.resize((max_px, round(im.height * max_px / im.width)), Image.LANCZOS)
    if im.mode == "RGBA":
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[3])
        im = bg
    im.save(dst, optimize=True)
    return dst


# Two cells of Phụ lục D and G were typed straight into Word with a plain space before
# "%"; the book keeps them that way, so the build leaves these two alone.
PLAIN_PERCENT = ("cùng phiên đã chậm đi 10–15 % do hạ xung",
                 "sau khi phát hiện hạ xung làm chậm 10–15 %")


def nbsp(line: str) -> str:
    """Tie a number to the percent sign after it so Word cannot break the line between."""
    line = re.sub(r"(\d) %", "\\1\u00a0%", line)
    for phrase in PLAIN_PERCENT:
        line = line.replace(phrase.replace(" %", "\u00a0%"), phrase)
    return line


def preprocess(md: str, work: Path) -> str:
    # covers are rebuilt natively in Word; pandoc gets everything from the abstract on
    md = md[md.index("# ABSTRACT"):]
    lines = md.split("\n")
    out: list[str] = []
    media_dir = work / "media"
    media_dir.mkdir(exist_ok=True)
    i = 0
    skip_until_break = None
    in_code = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
        elif not in_code:
            line = nbsp(line)

        # ---- regions whose literal content is replaced by a Word field ----------
        if skip_until_break:
            if stripped == PAGEBREAK_DIV:
                out.append(skip_until_break)
                out.append("")
                out.append("@@PAGEBREAK@@")
                skip_until_break = None
            i += 1
            continue

        m = re.match(r"^# (.+?)(?: \{[^}]*\})?\s*$", line)
        if m:
            title = m.group(1).strip()
            if title in FRONT_TITLES:
                # abstract pages: the thesis title sits above the page heading
                if title in ("ABSTRACT", "TÓM TẮT"):
                    j = i + 1
                    while not lines[j].strip():
                        j += 1
                    thesis_title = lines[j].strip().strip("*")
                    out += ['::: {custom-style="Abstract Title"}', thesis_title, ":::", ""]
                    lines[j] = ""
                out += ['::: {custom-style="Front Title"}', title, ":::", ""]
                if title == "MỤC LỤC":
                    skip_until_break = "@@TOC@@"
                i += 1
                continue
            if title == "DANH MỤC HÌNH ẢNH":
                out += [line, ""]
                skip_until_break = "@@TOF Hinh@@"
                i += 1
                continue
            if title == "DANH MỤC BẢNG BIỂU":
                out += [line, ""]
                skip_until_break = "@@TOF Bang@@"
                i += 1
                continue

        if stripped == PAGEBREAK_DIV:
            out += ["", "@@PAGEBREAK@@", ""]
            i += 1
            continue

        # ---- comment pages: dotted lines for the reviewer's handwriting ----------
        if re.fullmatch(r"\.{20,}", stripped):
            # collapse the run of dotted lines into a fixed number of ruled lines
            while i < len(lines) and (re.fullmatch(r"\.{20,}", lines[i].strip())
                                      or not lines[i].strip()):
                i += 1
            out += ["@@DOTS@@", ""] * 18        # the template rules 18 lines
            continue

        # ---- captions (chapters are numbered, appendices lettered) -----------------
        m = re.match(r"^\*((?:Hình|Bảng) (?:\d+|[A-Z])\.\d+ .+)\*\s*$", line)
        if m:
            out += ['::: {custom-style="Caption"}', m.group(1), ":::", ""]
            i += 1
            continue

        # ---- display equations: Word builds them later from the LaTeX source ------
        # The source travels hex-encoded inside a marker so that pandoc's Markdown
        # reader cannot touch a single backslash, underscore or brace of it.
        if stripped == "$$":
            j = i + 1
            body = []
            while lines[j].strip() != "$$":
                body.append(lines[j])
                j += 1
            tex = " ".join(b.strip() for b in body)
            tag = re.search(r"\\tag\{([^}]+)\}", tex)
            if not tag:
                sys.exit(f"công thức ở dòng {i + 1} thiếu \\tag{{...}}")
            tex = re.sub(r"\\tag\{[^}]+\}", "", tex)
            rows = [r.strip() for r in re.split(r"\\\\", tex) if r.strip()]
            encoded = " ".join(r.encode("utf-8").hex() for r in rows)
            out += ["", f"@@EQ {tag.group(1)} {encoded}@@", ""]
            i = j + 1
            continue

        # ---- images: compressed copy + explicit width ------------------------------
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)(\{[^}]*\})?\s*$", line)
        if m:
            src = (DOCS / m.group(2)).resolve()
            small = shrink_image(src, media_dir)
            width = 14.0 if src.name.startswith(NARROW_IMAGES) else TEXT_W_CM
            if src.stem in LANDSCAPE_IMAGES:
                # on its own landscape page the figure and its caption share the text
                # height (1,3 cm for the caption), so the width follows from the aspect
                # ratio, not from the margins
                im = Image.open(small)
                width = round(min(LAND_TEXT_W_CM,
                                  (LAND_TEXT_H_CM - 1.3) * im.width / im.height), 1)
                out += [f"@@LANDSCAPE {width}@@", ""]
            out.append(f"![{m.group(1)}]({small.as_posix()}){{width={width}cm}}")
            i += 1
            continue

        out.append(line)
        i += 1
    return "\n".join(out)


def run_pandoc(md_text: str, work: Path) -> Path:
    pre = work / "pre.md"
    pre.write_text(md_text, encoding="utf-8")
    raw = work / "raw.docx"
    cmd = ["pandoc", str(pre), "-f", "markdown-implicit_figures", "-t", "docx",
           "--syntax-highlighting=none", "-o", str(raw)]
    r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        sys.exit(f"pandoc thất bại:\n{r.stderr}")
    warnings = [w for w in r.stderr.splitlines() if w.startswith("[WARNING]")]
    for w in warnings:
        print("  pandoc:", w)
    return raw




# =============================================================================
# 2. Low-level OOXML helpers
# =============================================================================
def twips(cm: float) -> int:
    return round(cm / 2.54 * 1440)


def el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def child(parent, tag: str):
    """Existing child element, or a new one appended."""
    found = parent.find(qn(tag))
    if found is None:
        found = OxmlElement(tag)
        parent.append(found)
    return found


def set_fonts(rpr, font: str = FONT) -> None:
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:                      # rFonts must be the first child of rPr
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), font)


def add_field(paragraph, instr: str, cached: str = "", size_pt: float | None = None):
    """Append a complex field (begin / instr / separate / result / end)."""
    def run(child_el=None, text=None):
        r = paragraph.add_run()
        if size_pt:
            r.font.size = Pt(size_pt)
        if child_el is not None:
            r._r.append(child_el)
        if text is not None:
            r.text = text
        return r
    run(el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run(it)
    run(el("w:fldChar", **{"w:fldCharType": "separate"}))
    run(text=cached)
    run(el("w:fldChar", **{"w:fldCharType": "end"}))


# Child order required by the OOXML schema; Word rejects files that break it.
ORDER = {
    "w:pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
              "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
              "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
              "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
              "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"],
    "w:tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
                "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "w:trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
               "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"],
    "w:tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
               "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "w:sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type",
                 "pgSz", "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols",
                 "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi",
                 "rtlGutter", "docGrid", "printerSettings", "sectPrChange"],
}


def put(parent, new) -> None:
    """Insert new under parent at its schema position, replacing a same-named child."""
    def local(e):
        return e.tag.split("}")[1]
    key = "w:" + local(parent)
    seq = ORDER[key]
    for old in parent.findall(new.tag):
        parent.remove(old)
    rank = seq.index(local(new))
    for _i, c in enumerate(parent):
        name = local(c)
        if name in seq and seq.index(name) > rank:
            c.addprevious(new)
            return
    parent.append(new)


def para_border(paragraph, side: str) -> None:
    """The thin blue rule the template puts between header/footer and the text; the gap
    is measured so the rule lands where the template draws its line."""
    ppr = paragraph._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = el("w:pBdr")
        put(ppr, bdr)
    bdr.append(el(f"w:{side}", **{"w:val": "single", "w:sz": 4,
                                  "w:space": 13 if side == "bottom" else 9,
                                  "w:color": "5B9BD5"}))


def spacing(paragraph, before=0, after=0, line=1.5, align=None, left=None, first=None,
            keep_next=None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = Cm(left)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if keep_next is not None:
        pf.keep_with_next = keep_next


def clear_runs(paragraph) -> None:
    for c in list(paragraph._p):
        if c.tag != qn("w:pPr"):
            paragraph._p.remove(c)


def ensure_ppr(p_el):
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = el("w:pPr")
        p_el.insert(0, ppr)
    return ppr


# =============================================================================
# 3. Styles — one definition per rule of the template
# =============================================================================
def style_el(doc, style_id: str, name: str, based_on: str | None):
    styles = doc.styles.element
    for s in styles.findall(qn("w:style")):
        if s.get(qn("w:styleId")) == style_id:
            s.find(qn("w:name")).set(qn("w:val"), name)
            return s
    s = el("w:style", **{"w:type": "paragraph", "w:styleId": style_id})
    s.append(el("w:name", **{"w:val": name}))
    if based_on:
        s.append(el("w:basedOn", **{"w:val": based_on}))
    s.append(el("w:qFormat"))
    styles.append(s)
    return s


def define(doc, style_id, name, *, size=13, bold=False, italic=False, caps=False,
           align="left", before=0.0, after=0.0, line=1.5, left=0.0, first=0.0,
           keep_next=False, keep_lines=False, tabs=(), font=FONT, based_on="Normal",
           shade=None, outline=None):
    s = style_el(doc, style_id, name, based_on)
    for tag in ("w:pPr", "w:rPr"):
        old = s.find(qn(tag))
        if old is not None:
            s.remove(old)
    ppr = el("w:pPr")
    if keep_next:
        ppr.append(el("w:keepNext"))
    if keep_lines:
        ppr.append(el("w:keepLines"))
    if shade:
        ppr.append(el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": shade}))
    if tabs:
        t = el("w:tabs")
        for pos_cm, kind, leader in tabs:
            t.append(el("w:tab", **{"w:val": kind, "w:leader": leader, "w:pos": twips(pos_cm)}))
        ppr.append(t)
    ppr.append(el("w:spacing", **{"w:before": round(before * 20), "w:after": round(after * 20),
                                  "w:line": round(line * 240), "w:lineRule": "auto"}))
    ind = {"w:left": twips(left)}
    if first >= 0:
        ind["w:firstLine"] = twips(first)
    else:
        ind["w:hanging"] = twips(-first)
    ppr.append(el("w:ind", **ind))
    ppr.append(el("w:jc", **{"w:val": align}))
    if outline is not None:
        ppr.append(el("w:outlineLvl", **{"w:val": outline}))
    s.append(ppr)
    rpr = el("w:rPr")
    set_fonts(rpr, font)
    rpr.append(el("w:b", **({} if bold else {"w:val": "0"})))
    rpr.append(el("w:bCs", **({} if bold else {"w:val": "0"})))
    rpr.append(el("w:i", **({} if italic else {"w:val": "0"})))
    rpr.append(el("w:iCs", **({} if italic else {"w:val": "0"})))
    if caps:
        rpr.append(el("w:caps"))
    rpr.append(el("w:color", **{"w:val": "000000"}))
    rpr.append(el("w:sz", **{"w:val": round(size * 2)}))
    rpr.append(el("w:szCs", **{"w:val": round(size * 2)}))
    s.append(rpr)
    return s


def heading_numbering(doc) -> int:
    """A multilevel list like the template's: CHƯƠNG 1 / 1.1 / 1.1.1 / 1.1.1.1.

    The list is not tied to the heading styles, so a title that carries no number —
    "PHỤ LỤC", "Phụ lục A" — simply never gets a numPr and never moves the counters."""
    numbering = doc.part.numbering_part.element
    used_abs = {int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))}
    used_num = {int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))}
    abs_id, num_id = max(used_abs, default=0) + 1, max(used_num, default=0) + 1

    a = el("w:abstractNum", **{"w:abstractNumId": abs_id})
    a.append(el("w:multiLevelType", **{"w:val": "multilevel"}))
    # text at the same place as in the template: 1 cm, 1.27 cm, 1.52 cm; a chapter needs
    # room for "CHƯƠNG 1" itself, so its title starts at 3.5 cm
    levels = [("CHƯƠNG %1", 3.5), ("%1.%2", 1.0), ("%1.%2.%3", 1.27), ("%1.%2.%3.%4", 1.52)]
    for ilvl, (text, indent) in enumerate(levels):
        lvl = el("w:lvl", **{"w:ilvl": ilvl})
        lvl.append(el("w:start", **{"w:val": 1}))
        lvl.append(el("w:numFmt", **{"w:val": "decimal"}))
        lvl.append(el("w:lvlText", **{"w:val": text}))
        lvl.append(el("w:lvlJc", **{"w:val": "left"}))
        ppr = el("w:pPr")
        ppr.append(el("w:ind", **{"w:left": twips(indent), "w:hanging": twips(indent)}))
        lvl.append(ppr)
        a.append(lvl)
    numbering.insert(0, a)
    n = el("w:num", **{"w:numId": num_id})
    n.append(el("w:abstractNumId", **{"w:val": abs_id}))
    numbering.append(n)
    return num_id


def list_markers(doc) -> int:
    """Give every list the markers the template uses: an ordered list reads "1)", "2)"
    rather than "1.", "2.", and an unordered one is marked with a hyphen, not a dot.

    Pandoc writes the Word default — U+F0B7 in the Symbol font, a round dot — at every
    bullet level. The marker is set on the numbering level, so it also has to carry its
    own font: leaving Symbol behind would print the hyphen from that font instead of the
    Times New Roman one."""
    changed = 0
    for lvl in doc.part.numbering_part.element.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        text = lvl.find(qn("w:lvlText"))
        if fmt is None or text is None:
            continue
        kind, val = fmt.get(qn("w:val")), text.get(qn("w:val"))
        if kind == "decimal" and re.fullmatch(r"%\d\.", val):
            text.set(qn("w:val"), val[:-1] + ")")
            changed += 1
        elif kind == "bullet":
            text.set(qn("w:val"), BULLET)
            rpr = child(lvl, "w:rPr")       # rPr is the last child of w:lvl
            set_fonts(rpr)
            rpr.find(qn("w:rFonts")).attrib.pop(qn("w:hint"), None)
            changed += 1
    return changed


def number_heading(p, num_id: int, ilvl: int) -> None:
    npr = el("w:numPr")
    npr.append(el("w:ilvl", **{"w:val": ilvl}))
    npr.append(el("w:numId", **{"w:val": num_id}))
    put(ensure_ppr(p._p), npr)


def apply_styles(doc) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_fonts(rpr)
    lang = child(rpr, "w:lang")
    lang.set(qn("w:val"), "vi-VN")
    lang.set(qn("w:eastAsia"), "vi-VN")

    dots = [(TEXT_W_CM, "right", "dot")]
    # "Các danh mục … Normal: 13, before 0, after 0, 1.15 lines"
    define(doc, "Normal", "Normal", line=1.15, based_on=None)
    # "Content: 13, before 10 pt, after 0, 1.5 lines, không thụt đầu dòng, căn hai bên".
    # pandoc's body styles are renamed after the template's own style names.
    define(doc, "BodyText", "Content", before=10, line=1.5, align="both")
    # pandoc gives the paragraph right after a heading a style of its own; the template
    # knows a single body style, so every one of them is pointed at Content
    for st in doc.element.body.iter(qn("w:pStyle")):
        if st.get(qn("w:val")) == "FirstParagraph":
            st.set(qn("w:val"), "BodyText")
    # "Bullet: 13, 0/0, 1.5 lines, canh trái" — the indent is set on each list item,
    # because table cells share this style
    define(doc, "Compact", "Bullet", line=1.5)
    # "Heading 1: 14, đậm, in hoa, 24/24, 1.15 lines, căn trái"; the hanging indent leaves
    # "CHƯƠNG n" and the tab after it their own column, as in the template
    define(doc, "Heading1", "heading 1", size=14, bold=True, caps=True, before=24, after=24,
           line=1.15, keep_next=True, outline=0, left=3.5, first=-3.5)
    # "Heading 2: 13, đậm và đứng, 6/12, 1.15"
    define(doc, "Heading2", "heading 2", bold=True, before=6, after=12, line=1.15,
           keep_next=True, outline=1, left=1.0, first=-1.0)
    # "Heading 3: 13, đậm và nghiêng, 6/12, 1.15"
    define(doc, "Heading3", "heading 3", bold=True, italic=True, before=6, after=12,
           line=1.15, keep_next=True, outline=2, left=1.27, first=-1.27)
    # "Heading 4: 13, nghiêng, 6/12, 1.15"
    define(doc, "Heading4", "heading 4", italic=True, before=6, after=12, line=1.15,
           keep_next=True, outline=3, left=1.52, first=-1.52)
    # "Caption: 13, 6/12, 1.15, căn giữa" — italic, the way the template's own style is
    define(doc, "Caption", "caption", italic=True, before=6, after=12, line=1.15,
           align="center")
    # "Heading 1N" of the template: a chapter-sized title with no number, centred, still
    # picked up by the table of contents — but at outline level 5, so that its entries
    # land in "toc 5" instead of "toc 1". A numbered chapter entry and an unnumbered page
    # title need opposite indents (see TOC1 below), and one TOC level can only carry one
    # style. The price is cosmetic and confined to Word itself: the navigation pane shows
    # these titles nested rather than at the top. Word's other way of splitting them, the
    # TOC field's \t switch, was measured not to match this style name at all.
    define(doc, "Heading1N", "Heading 1N", size=14, bold=True, caps=True, before=24,
           after=24, line=1.15, align="center", keep_next=True, outline=4)
    # front-matter page titles: the template's CoverB, and out of the table of contents
    define(doc, "FrontTitle", "CoverB", size=14, bold=True, caps=True, after=12,
           line=1.15, align="center", keep_next=True)
    # the template sets the thesis title above ABSTRACT / TÓM TẮT in 14 pt bold, flush left
    define(doc, "AbstractTitle", "Abstract Title", size=14, bold=True, after=12, line=1.5,
           align="left")
    define(doc, "BlockText", "Block Text", line=1.5, align="right")
    # a code block is one paragraph (lines joined by breaks); like a table, it moves to
    # the next page whole rather than split
    define(doc, "SourceCode", "Source Code", size=10, line=1.0, font="Consolas",
           shade="F2F2F2", keep_lines=True)
    # TOC and lists of figures/tables: 13, 0/0 pt, 1.15 lines, dotted right tab; the
    # left indents per level are the template's
    # Chapter 2 does not describe the table of contents, so it follows the template: the
    # entries run at 1,3 lines with 5 pt after, and a sub-section starts 1 cm further in
    # than its indent, putting its number at 1,46 cm (level 2) and 1,92 cm (level 3).
    # a chapter entry reads "CHƯƠNG 4 <tab> TITLE <tab> 44"; Word's TOC field writes that
    # first tab at 1.680 twips. Without an indent, a title long enough to wrap puts its
    # tail back at 0 and the tab after it finds that 1.680 stop instead of the dotted
    # right one, dropping the page number mid-line with no leader. Hanging the same
    # amount keeps the first line where it was and starts the tail under the title.
    define(doc, "TOC1", "toc 1", after=5, line=1.3, tabs=dots,
           left=TOC1_HANG_CM, first=-TOC1_HANG_CM)
    define(doc, "TOC2", "toc 2", after=5, left=0.46, first=1.0, line=1.3, tabs=dots)
    define(doc, "TOC3", "toc 3", after=5, left=0.92, first=1.0, line=1.3, tabs=dots)
    # level 4 keeps the ladder going, for a sub-section numbered with four digits
    define(doc, "TOC4", "toc 4", after=5, left=1.38, first=1.0, line=1.3, tabs=dots)
    # level 5 belongs to the unnumbered page titles — "PHỤ LỤC", "TÀI LIỆU THAM KHẢO",
    # the lists. They carry no number, so they need no column for one, and above all no
    # hanging indent: that indent is itself a tab stop, and "PHỤ LỤC" is short enough to
    # end before it, which used to send the tab before the page number to 2,96 cm and
    # print "PHỤ LỤC     g" with no dot leader. Flush left, the only stop left is the
    # dotted one at the right margin, which is where the page number belongs.
    define(doc, "TOC5", "toc 5", after=5, line=1.3, tabs=dots)
    # the lists of figures and tables do have a rule: "Normal, 13, 0/0 pt, 1.15 lines"
    define(doc, "TableofFigures", "table of figures", line=1.15, tabs=dots)

    for s in doc.styles.element.findall(qn("w:style")):
        sid = s.get(qn("w:styleId"))
        if sid in ("VerbatimChar", "Hyperlink"):
            rpr = child(s, "w:rPr")
            for c in list(rpr):
                rpr.remove(c)
            if sid == "VerbatimChar":
                set_fonts(rpr, "Consolas")
                rpr.append(el("w:sz", **{"w:val": 22}))
            else:
                rpr.append(el("w:color", **{"w:val": "000000"}))


# =============================================================================
# 4. Body pass — markers, captions, images, equations, lists, references
# =============================================================================
def block_text(e) -> str:
    return "".join(t.text or "" for t in e.iter(qn("w:t"))) if e is not None else ""


def next_block(p):
    e = p._p.getnext()
    while e is not None and e.tag not in (qn("w:p"), qn("w:tbl")):
        e = e.getnext()
    return e


EQ_SIDE_CM = 1.5            # number column; the empty left column mirrors it


def no_borders(parent, tag: str) -> None:
    """A tblBorders / tcBorders element with every side switched off."""
    borders = el(tag)
    inner = ("insideH", "insideV") if tag == "w:tblBorders" else ()
    sides = ("top", "left", "bottom", "right") + inner
    for side in sides:
        borders.append(el(f"w:{side}", **{"w:val": "nil"}))
    put(parent, borders)


def make_equation_table(doc, marker, rows: list[str], chapter: str, number: str,
                        first_in_chapter: bool) -> None:
    """Replace the marker paragraph by an invisible 1 x 3 table: nothing, the LaTeX source
    of the equation (one paragraph per line, turned into OMML by Word later), and the
    number "(chapter.n)" whose n is a SEQ field, so Word renumbers when equations move."""
    table = doc.add_table(rows=1, cols=3)
    tbl = table._tbl
    marker._p.addprevious(tbl)
    tblpr = tbl.tblPr
    for old in tblpr.findall(qn("w:tblStyle")):
        tblpr.remove(old)
    no_borders(tblpr, "w:tblBorders")
    put(tblpr, el("w:tblW", **{"w:w": twips(TEXT_W_CM), "w:type": "dxa"}))
    put(tblpr, el("w:tblLayout", **{"w:type": "fixed"}))
    put(tblpr, el("w:tblLook", **{"w:val": "0000", "w:firstRow": "0", "w:lastRow": "0",
                                  "w:firstColumn": "0", "w:lastColumn": "0",
                                  "w:noHBand": "1", "w:noVBand": "1"}))
    put(tblpr, el("w:tblCaption", **{"w:val": "EQ"}))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (EQ_SIDE_CM, TEXT_W_CM - 2 * EQ_SIDE_CM, EQ_SIDE_CM)
    for col, w in zip(table.columns, widths, strict=False):
        col.width = Cm(w)
    for cell, w in zip(table.rows[0].cells, widths, strict=False):
        cell.width = Cm(w)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        no_borders(cell._tc.get_or_add_tcPr(), "w:tcBorders")
    put(table.rows[0]._tr.get_or_add_trPr(), el("w:cantSplit"))

    side = table.rows[0].cells[0].paragraphs[0]
    spacing(side, line=1.0)
    mid = table.rows[0].cells[1]
    for k, tex in enumerate(rows):
        par = mid.paragraphs[0] if k == 0 else mid.add_paragraph()
        par.style = doc.styles["Normal"]
        par.add_run(tex)
        spacing(par, before=6 if k == 0 else 3, after=6 if k == len(rows) - 1 else 3,
                line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    num = table.rows[0].cells[2].paragraphs[0]
    num.add_run(f"({chapter}.")
    add_field(num, "SEQ CongThuc \\* ARABIC" + (" \\r 1" if first_in_chapter else ""), number)
    num.add_run(")")
    spacing(num, line=1.0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    marker._p.getparent().remove(marker._p)


PANDOC_CAP_CM = 14.82        # the widest picture pandoc writes, whatever {width=} says


def picture_width_cm(p) -> float:
    """Width of the first inline picture in p."""
    ext = next(p._p.iter(qn("wp:inline"))).find(qn("wp:extent"))
    return int(ext.get("cx")) / 360000


def set_picture_width(p, width_cm: float) -> None:
    """Resize every inline picture in p to width_cm, keeping its aspect ratio.

    pandoc caps images at its own default text width, so a figure meant for a
    landscape page has to be widened after the fact."""
    emu = round(width_cm * 360000)
    for inline in p._p.iter(qn("wp:inline")):
        ext = inline.find(qn("wp:extent"))
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        new_cy = round(cy * emu / cx)
        ext.set("cx", str(emu))
        ext.set("cy", str(new_cy))
        for a_ext in inline.iter(qn("a:ext")):
            if a_ext.get("cx") is not None:
                a_ext.set("cx", str(emu))
                a_ext.set("cy", str(new_cy))


def rebuild_caption(p, first_in_chapter: bool) -> str:
    m = re.match(r"^(Hình|Bảng) (\d+|[A-Z])\.(\d+) (.+)$", p.text.strip())
    if not m:
        return ""
    kind, ch, n, rest = m.groups()
    seq = "Hinh" if kind == "Hình" else "Bang"
    clear_runs(p)
    p.add_run(f"{kind} {ch}.")
    add_field(p, f"SEQ {seq} \\* ARABIC" + (" \\r 1" if first_in_chapter else ""), n)
    p.add_run(f" {rest}")
    return kind


def is_bullet(doc, numpr) -> bool:
    """True when this list item is a bullet rather than a number."""
    num_id = numpr.find(qn("w:numId")).get(qn("w:val"))
    ilvl = numpr.find(qn("w:ilvl"))
    want = ilvl.get(qn("w:val")) if ilvl is not None else "0"
    numbering = doc.part.numbering_part.element
    abstract = None
    for n in numbering.findall(qn("w:num")):
        if n.get(qn("w:numId")) == num_id:
            abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == abstract:
            for lvl in a.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == want:
                    fmt = lvl.find(qn("w:numFmt"))
                    return fmt is not None and fmt.get(qn("w:val")) == "bullet"
    return False


def strip_prefix(p, prefix: str) -> None:
    """Drop the literal number in front of a heading — Word draws it from the list."""
    for r in p.runs:
        if r.text.startswith(prefix):
            r.text = r.text[len(prefix):]
            return
    rest = p.text[len(prefix):]
    clear_runs(p)
    p.add_run(rest)


def body_pass(doc, num_id: int) -> dict:
    counts = {"trang": 0, "phần": 0, "công thức": 0, "dòng công thức": 0, "chú thích": 0,
              "ảnh": 0, "đề mục đánh số": 0}
    seen_kind: set[str] = set()
    section_ends = []            # (paragraph, role) in document order
    in_refs = False
    landscape_open = False
    landscape_width = LAND_TEXT_W_CM
    counters = [0, 0, 0, 0]      # what Word's list will count, checked against the draft
    for p in list(doc.paragraphs):
        if p._p.getparent() is None:
            continue
        text = p.text.strip()
        style = p.style.name.lower()

        if style == "heading 1":
            seen_kind = set()
            in_refs = text.startswith("TÀI LIỆU THAM KHẢO")
        elif style == "heading 2" and re.match(r"Phụ lục [A-Z]\b", text):
            seen_kind = set()            # "Hình A.1", "Hình B.1": each appendix restarts

        # ---- heading numbers: Word's multilevel list replaces the literal ones --------
        if style in ("heading 1", "heading 2", "heading 3", "heading 4"):
            lvl = int(style[-1]) - 1
            m = re.match(r"^CHƯƠNG (\d+) " if lvl == 0 else r"^(\d+(?:\.\d+)+) ", text)
            if lvl == 0 and text.startswith(UNNUMBERED_TITLES):
                p.style = doc.styles["Heading 1N"]
                spacing(p, before=24, after=24, line=1.15, left=0, first=0, keep_next=True)
            elif m:
                counters[lvl] += 1
                for k in range(lvl + 1, 4):
                    counters[k] = 0
                want = ".".join(str(counters[k]) for k in range(lvl + 1))
                if m.group(1) != want:
                    sys.exit(f"số đề mục trong bản thảo là {m.group(1)} nhưng Word sẽ đánh "
                             f"số {want}: {text[:60]}")
                strip_prefix(p, m.group(0))
                number_heading(p, num_id, lvl)
                counts["đề mục đánh số"] += 1
            else:                        # "Phụ lục A — …": a title that carries no number
                spacing(p, before=6, after=12, line=1.15, left=0, first=0, keep_next=True)
            continue

        if text == "@@PAGEBREAK@@":
            nxt = next_block(p)
            nxt_text = block_text(nxt)
            if nxt_text.startswith("CHƯƠNG 1 ") or nxt_text.startswith("TÀI LIỆU THAM KHẢO"):
                clear_runs(p)
                spacing(p, line=1.0)
                rpr = el("w:rPr")
                rpr.append(el("w:sz", **{"w:val": 2}))
                put(ensure_ppr(p._p), rpr)
                section_ends.append((p, "front" if nxt_text.startswith("CHƯƠNG 1 ") else "main"))
                counts["phần"] += 1
            else:
                if nxt is not None and nxt.tag == qn("w:p"):
                    put(ensure_ppr(nxt), el("w:pageBreakBefore"))
                p._p.getparent().remove(p._p)
                counts["trang"] += 1
            continue

        m = re.fullmatch(r"@@LANDSCAPE ([\d.]+)@@", text)
        if m:
            landscape_width = float(m.group(1))
            clear_runs(p)
            spacing(p, line=1.0)
            rpr = el("w:rPr")
            rpr.append(el("w:sz", **{"w:val": 2}))
            put(ensure_ppr(p._p), rpr)
            section_ends.append((p, "portrait"))
            landscape_open = True
            continue

        if text == "@@TOC@@":
            clear_runs(p)
            # levels 1–4 are the numbered headings, level 5 the unnumbered page titles
            add_field(p, 'TOC \\o "1-5" \\h \\z \\u', "Nhấn F9 để cập nhật mục lục")
            continue
        m = re.match(r"@@TOF (\w+)@@", text)
        if m:
            clear_runs(p)
            add_field(p, f'TOC \\h \\z \\c "{m.group(1)}"', "Nhấn F9 để cập nhật danh mục")
            continue

        if text == "@@DOTS@@":
            clear_runs(p)
            p.add_run("\t")
            spacing(p, before=6, line=1.5)      # the template's ruled lines, 28 pt apart
            p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)
            continue

        m = re.fullmatch(r"@@EQ (\d+)\.(\d+) ([0-9a-f ]+)@@", text)
        if m:
            chapter, number, encoded = m.groups()
            rows = [bytes.fromhex(h).decode("utf-8") for h in encoded.split()]
            make_equation_table(doc, p, rows, chapter, number,
                                first_in_chapter="CongThuc" not in seen_kind)
            seen_kind.add("CongThuc")
            counts["công thức"] += 1
            counts["dòng công thức"] += len(rows)
            continue
        assert not text.startswith("@@EQ"), f"dấu công thức hỏng: {text[:60]}"

        if style == "caption":
            kind = rebuild_caption(p, first_in_chapter=text.split(" ")[0] not in seen_kind)
            if kind:
                seen_kind.add(kind)
                counts["chú thích"] += 1
                if kind == "Hình" and landscape_open:
                    section_ends.append((p, "landscape"))
                    landscape_open = False
                if kind == "Bảng":
                    p.paragraph_format.keep_with_next = True
                else:
                    prev = p._p.getprevious()
                    if prev is not None and prev.tag == qn("w:p"):
                        put(ensure_ppr(prev), el("w:keepNext"))
            continue

        if p._p.findall(".//" + qn("w:drawing")):
            spacing(p, before=6, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
            if landscape_open:
                set_picture_width(p, landscape_width)
            elif picture_width_cm(p) > PANDOC_CAP_CM - 0.05:
                # pandoc caps every picture at 14,82 cm — its own default text width —
                # so a figure asked for at 15,5 cm came out 5 % smaller, text and all
                set_picture_width(p, TEXT_W_CM)
            counts["ảnh"] += 1
            continue

        if text.startswith("TP. Hồ Chí Minh, ngày"):
            spacing(p, before=18, line=1.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
            continue
        if text == "CHỮ KÝ CỦA GIẢNG VIÊN":
            spacing(p, line=1.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
            continue                           # the template leaves this line unbolded

        if in_refs and re.match(r"^\[\d+\]\. ", text):
            # "những dòng tiếp sau cần phải lùi vào 1 cm so với dòng đầu tiên"
            spacing(p, before=6, line=1.5, align=WD_ALIGN_PARAGRAPH.LEFT, left=1.0, first=-1.0)
            continue

        # a lead-in set in bold on a line of its own ("Khó khăn.") stays on the page of the
        # list it opens, instead of ending a page alone
        nxt = p._p.getnext()
        runs = [r for r in p.runs if r.text.strip()]
        if (runs and (p._p.pPr is None or p._p.pPr.numPr is None)
                and all(r.bold or r.style.name == "Strong" for r in runs)
                and nxt is not None and nxt.tag == qn("w:p")
                and nxt.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None):
            p.paragraph_format.keep_with_next = True

        if p._p.pPr is not None and p._p.pPr.numPr is not None:
            npr = p._p.pPr.numPr
            ilvl = npr.find(qn("w:ilvl"))
            depth = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
            # the template keeps a bullet at the margin but sets a numbered list in,
            # its number at 1 cm and the text at 1,6 cm
            numbered = npr.find(qn("w:numId")) is not None and not is_bullet(doc, npr)
            left = (1.6 if numbered else 0.5) + 0.75 * depth
            spacing(p, line=1.5, align=WD_ALIGN_PARAGRAPH.LEFT,
                    left=left, first=-0.6 if numbered else -0.5)
    counts["section_ends"] = section_ends
    return counts


# =============================================================================
# 5. Tables
# =============================================================================
def cell_text(cell) -> str:
    """Visible text of a cell, the characters of its inline equations included."""
    return "".join(t.text or "" for t in cell._tc.iter(qn("w:t"), qn("m:t")))


def fit_columns(table, width_cm: float, pt: float) -> None:
    """Column widths from content: each column gets at least its longest word, and the
    rest of the text width goes to columns in proportion to how much text they carry."""
    char_cm = 0.02 * pt                       # a generous Times New Roman glyph width
    cols = list(zip(*[[cell_text(c) for c in row.cells] for row in table.rows], strict=False))
    dense = len(cols) >= 8
    # cell padding: Word's default 0.19 cm a side; dense tables get 0.1 cm
    pad = 0.25 if dense else 0.45
    if dense:
        mar = el("w:tblCellMar")
        for side in ("left", "right"):
            mar.append(el(f"w:{side}", **{"w:w": twips(0.1), "w:type": "dxa"}))
        put(table._tbl.tblPr, mar)
    mins, wants = [], []
    for texts in cols:
        # Word may break after a space, an en dash or a slash, never inside a number
        pieces = [w for t in texts for w in re.split(r"[\s–/]+", t) if w]
        longest_word = max((len(w) for w in pieces), default=1)
        lens = [len(t) for t in texts[1:]] or [len(texts[0])]
        # the widest word decides the column: it is often set bold, or in Cambria Math,
        # both wider than the plain face char_cm is measured on
        mins.append(longest_word * char_cm * 1.3 + pad)
        wants.append(max(sum(lens) / len(lens), 0.5 * max(lens), len(texts[0]) * 0.6))
    spare = width_cm - sum(mins)
    if spare <= 0:
        widths = [m * width_cm / sum(mins) for m in mins]
    else:
        extra = [max(w * char_cm - m, 0.05) for w, m in zip(wants, mins, strict=False)]
        widths = [m + spare * e / sum(extra) for m, e in zip(mins, extra, strict=False)]
    tblpr = table._tbl.tblPr
    put(tblpr, el("w:tblW", **{"w:w": twips(width_cm), "w:type": "dxa"}))
    put(tblpr, el("w:tblLayout", **{"w:type": "fixed"}))
    for col, w in zip(table.columns, widths, strict=False):
        col.width = Cm(w)
    for row in table.rows:
        for cell, w in zip(row.cells, widths, strict=False):
            cell.width = Cm(w)


# the two front-matter lists are laid out like the template's own: 13 pt, no shading,
# a plain grid, the term centred and bold, its explanation on the left
GLOSSARY_HEADS = {("Từ viết tắt", "Từ đầy đủ", "Nghĩa"): (3.14, 8.03, 4.30),
                  ("Ký hiệu", "Ý nghĩa"): (3.14, TEXT_W_CM - 3.14)}

# Seven data tables whose columns were set by hand in Word for the submitted book (v4):
# Bảng 7.2, D.1–D.4, G.1 and G.2. Their widths are copied from it, to the twip, so the
# build lays them out the same way; every other data table keeps fit_columns().
BOOK_WIDTHS = {
    ("Cấu hình", "Tham số — lookback gộp", "Tham số — chunk-aware", "ECE — lookback gộp",
     "ECE — chunk-aware", "Bên có ECE thấp hơn"): (2.999, 2.142, 2.434, 2.140, 2.434, 3.351),
    ("Thành phần", "Cách cố định"): (4.233, 11.266),
    ("Môi trường", "Python", "PyTorch", "transformers", "Dùng cho"):
        (3.351, 2.028, 2.646, 2.822, 4.652),
    ("Loại kết quả", "Mức tái lập", "Ghi chú"): (4.057, 3.704, 7.738),
    ("Tệp", "SHA-256"): (4.674, 10.825),
    ("Hạng mục", "Nhóm thực hiện và quyết định", "Công cụ AI hỗ trợ"): (3.704, 6.703, 5.092),
    ("Ngày", "Yêu cầu (trích)", "Mục đích"): (1.764, 7.585, 6.151),
}
# the hash table (Bảng D.4) sets its rows at 8 pt so every SHA-256 fits on one line, as
# the book does; its header row keeps the usual size
BODY_PT = {("Tệp", "SHA-256"): 8}


def set_widths(table, widths) -> None:
    """Fix the column widths instead of letting Word fit them to the content."""
    tblpr = table._tbl.tblPr
    put(tblpr, el("w:tblW", **{"w:w": twips(sum(widths)), "w:type": "dxa"}))
    put(tblpr, el("w:tblLayout", **{"w:type": "fixed"}))
    for col, w in zip(table.columns, widths, strict=False):
        col.width = Cm(w)
    for row in table.rows:
        for cell, w in zip(row.cells, widths, strict=False):
            cell.width = Cm(w)


def table_height_cm(table, pt: float) -> float:
    """Rough laid-out height of a table whose column widths are set: every row is as tall
    as its fullest cell, a line of text is 1.15 × 1.15 × pt (Times New Roman's line at
    1,15 spacing) and each row adds 4 pt of paragraph spacing. It only has to tell a table
    that fits on a page from one that cannot, and the two kinds are far apart."""
    line_cm = pt * 1.15 * 1.15 / 28.35
    char_cm = 0.018 * pt                      # an average glyph, not fit_columns' widest
    total = 0.0
    for row in table.rows:
        lines = 1
        for cell in row.cells:
            usable = max(cell.width.cm - 0.4, 0.5)
            lines = max(lines, sum(max(1, math.ceil(len(cp.text) * char_cm / usable))
                                   for cp in cell.paragraphs))
        total += lines * line_cm + 4 / 28.35
    return total


def format_tables(doc) -> int:
    n = 0
    long_tables = []
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        n += 1
        head = [cell_text(c).strip() for c in table.rows[0].cells]
        is_log = " ".join(head).startswith("Tuần")
        is_glossary = tuple(head) in GLOSSARY_HEADS
        n_cols = len(table.columns)
        # template allows 10–13 pt; the glossaries follow the template's own 13 pt
        pt = 13 if is_glossary else (10 if n_cols >= 8 else TABLE_PT)
        last = len(table.rows) - 1
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        borders = el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(el(f"w:{side}", **{"w:val": "single", "w:sz": 4, "w:space": 0,
                                              "w:color": "000000"}))
        put(tblpr, borders)
        put(tblpr, el("w:tblCaption",
                      **{"w:val": "LOG" if is_log else "GLOSSARY" if is_glossary else "DATA"}))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            put(trpr, el("w:cantSplit"))
            if ri == 0:
                put(trpr, el("w:tblHeader"))
            elif is_log:
                put(trpr, el("w:trHeight", **{"w:val": twips(1.3), "w:hRule": "atLeast"}))
            for ci, cell in enumerate(row.cells):
                if ri == 0:              # no fill anywhere: the template's tables are plain
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                centred = ri == 0 or (is_glossary and ci == 0)
                for cp in cell.paragraphs:
                    spacing(cp, before=2, after=2, line=1.15, left=0, first=0,
                            align=WD_ALIGN_PARAGRAPH.CENTER if centred
                            else WD_ALIGN_PARAGRAPH.LEFT)
                    for r in cp.runs:
                        r.font.size = Pt(pt if ri == 0 else BODY_PT.get(tuple(head), pt))
        if is_log:
            # the two date columns must hold "03/08/2026" on one line
            mar = el("w:tblCellMar")
            for side in ("left", "right"):
                mar.append(el(f"w:{side}", **{"w:w": twips(0.12), "w:type": "dxa"}))
            put(tblpr, mar)
            set_widths(table, (1.2, 2.4, 2.4, 1.9, 4.8, 2.8))
        elif is_glossary:
            set_widths(table, GLOSSARY_HEADS[tuple(head)])
        elif tuple(head) in BOOK_WIDTHS:
            set_widths(table, BOOK_WIDTHS[tuple(head)])
        else:
            fit_columns(table, TEXT_W_CM, pt)
        # "các hình, bảng biểu phải được trình bày trong một trang (ngoại trừ các bảng dài
        # hơn một trang)": a table that fits on a page with its caption is chained row to
        # row, so Word moves the whole of it to the next page rather than split it. Only
        # the longer ones may break, and their header row repeats.
        height = table_height_cm(table, pt)
        if not is_log and height + CAPTION_CM <= TEXT_H_CM:
            for row in table.rows[:last]:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        else:
            long_tables.append(f"{head[0][:24]} ({height:.0f} cm)")
    print(f"   {n - len(long_tables)} bảng giữ trọn một trang; được ngắt vì dài hơn một "
          f"trang: {', '.join(long_tables) or 'không có'}")
    return n


# =============================================================================
# 6. Covers, sections, headers and footers
# =============================================================================
def insert_covers(doc):
    """Two identical Vietnamese covers (outer + inner) and the English one.

    Line sizes and the vertical rhythm are measured off the template's own covers: the
    faculty block at the top, the logo, the students, the title at 20 pt, the major, the
    supervisor and the city at the foot."""
    anchor = doc.paragraphs[0]._p
    ends = []

    def line(text, size=14, bold=True, before=0.0, line_sp=1.0, picture=None):
        p = doc.add_paragraph()
        anchor.addprevious(p._p)
        p.style = doc.styles["Normal"]
        spacing(p, before=before, line=line_sp, align=WD_ALIGN_PARAGRAPH.CENTER,
                left=0, first=0)
        if picture:
            # width only: the height follows the logo's own 3,51 : 1. The .docx keeps the
            # file's 1.600 px; Word's PDF export resamples every picture to ~200 ppi
            # whatever the document settings (tested 27/09), so the PDF shows 501 px
            p.add_run().add_picture(str(picture), width=Inches(2.51))
        else:
            r = p.add_run(text)
            r.font.size, r.font.bold = Pt(size), bold
        return p

    def cover(top, students, title, major, supervisor, city, city_size):
        for k, t in enumerate(top):
            line(t, bold=not t.startswith("BỘ CÔNG THƯƠNG"), before=30 if k == 0 else 0)
        line("", picture=LOGO, before=50)
        line(students[0], before=38)
        line(students[1], before=11)
        line(title, size=20, before=40)
        line(major, before=55)
        line(supervisor, before=55)
        ends.append(line(city, size=city_size, bold=False, before=60))

    vi = (["BỘ CÔNG THƯƠNG", "TRƯỜNG ĐẠI HỌC CÔNG NGHIỆP TP. HỒ CHÍ MINH",
           "KHOA CÔNG NGHỆ THÔNG TIN"], STUDENTS_VI, TITLE_VI, "Ngành: Khoa học dữ liệu",
          "Giảng viên hướng dẫn: ThS. Trương Vĩnh Linh",
          "THÀNH PHỐ HỒ CHÍ MINH, THÁNG 12 NĂM 2026", 12)
    cover(*vi)
    cover(*vi)
    cover(["INDUSTRIAL UNIVERSITY OF HO CHI MINH CITY", "FACULTY OF INFORMATION TECHNOLOGY"],
          STUDENTS_EN, TITLE_EN, "Major: Data Science", "Supervisor: MSc. Truong Vinh Linh",
          "HO CHI MINH CITY – DECEMBER 2026", 14)
    return ends


def page_border(sectpr, kind: str) -> None:
    """The template frames the outer Vietnamese cover with Word's twisted-lines art and
    the English one with a plain 1 pt rectangle; every other page is left bare.

    Word measures the offsets from the text — but only while the section owns no header
    or footer, otherwise it wraps the frame around those instead and the frame drifts to
    the paper edge. The cover sections therefore stay without one (see word_finish)."""
    for old in sectpr.findall(qn("w:pgBorders")):
        sectpr.remove(old)
    if kind == "none":
        return
    art = kind == "art"
    borders = el("w:pgBorders")
    for side in ("top", "left", "bottom", "right"):
        wide = side in ("left", "right")
        borders.append(el(f"w:{side}", **{
            "w:val": "twistedLines1" if art else "single",
            "w:sz": 18 if art else 8,
            "w:space": (4 if wide else 1) if art else (19 if wide else 18),
            "w:color": "auto"}))
    put(sectpr, borders)


def set_pgnum(sectpr, fmt: str, start: int | None) -> None:
    for old in sectpr.findall(qn("w:pgNumType")):
        sectpr.remove(old)
    pg = el("w:pgNumType", **{"w:fmt": fmt})
    if start is not None:
        pg.set(qn("w:start"), str(start))
    # schema order: … pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols …
    anchor = sectpr.find(qn("w:cols"))
    if anchor is not None:
        anchor.addprevious(pg)
    else:
        sectpr.append(pg)


def page_header(section) -> None:
    """The thesis in the header, 13 pt and flush left as in the template, with the same
    blue rule under it.

    Given to the first chapter section only — the front matter before it has none, and
    neither may the covers: a section that owns a header has its page border drawn
    around that header instead of around the text, which would push the frame on the
    covers out to the paper edge. Later sections inherit it."""
    part = section.header
    part.is_linked_to_previous = False
    h = part.paragraphs[0]
    clear_runs(h)
    spacing(h, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, left=0, first=0)
    r = h.add_run(HEADER_TEXT)
    r.font.size = Pt(13)
    para_border(h, "bottom")


def page_footer(section, width_cm: float) -> None:
    """Students + page number, 13 pt as in the template, separated from the text by a
    rule; the number sits bottom right. Like the header, never given to a cover section."""
    part = section.footer
    part.is_linked_to_previous = False
    f = part.paragraphs[0]
    clear_runs(f)
    spacing(f, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, left=0, first=0)
    f.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)
    r = f.add_run(FOOTER_TEXT + "\t")
    r.font.size = Pt(13)
    add_field(f, "PAGE", "1", size_pt=13)
    para_border(f, "top")


def build_sections(doc, cover_ends, section_ends) -> None:
    base = doc.sections[-1]
    base.page_width, base.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
    base.top_margin, base.bottom_margin = Cm(MARGIN["top"]), Cm(MARGIN["bottom"])
    base.left_margin, base.right_margin = Cm(MARGIN["left"]), Cm(MARGIN["right"])
    base.header_distance, base.footer_distance = Pt(50), Pt(40)       # as the template
    base.start_type = WD_SECTION.NEW_PAGE
    body_sectpr = doc.element.body.find(qn("w:sectPr"))
    for ref in (body_sectpr.findall(qn("w:headerReference"))
                + body_sectpr.findall(qn("w:footerReference"))):
        body_sectpr.remove(ref)

    ends = [(p, f"cover{k + 1}") for k, p in enumerate(cover_ends)] + section_ends
    for p, _ in ends:
        ensure_ppr(p._p).append(copy.deepcopy(body_sectpr))
    roles = [r for _, r in ends] + ["back"]
    sections = list(doc.sections)
    assert len(sections) == len(roles), (len(sections), roles)

    first_main = True
    previous_width = None
    for sec, role in zip(sections, roles, strict=False):
        if role.startswith("cover"):
            set_pgnum(sec._sectPr, "lowerRoman", 1 if role == "cover1" else None)
            page_border(sec._sectPr, {"cover1": "art", "cover3": "plain"}.get(role, "none"))
            continue
        if role == "front":
            # keeps counting from the covers; the footer starts here, the header only
            # with chapter 1 (the user's call, 28/09)
            set_pgnum(sec._sectPr, "lowerRoman", None)
        elif role in ("portrait", "landscape", "main"):
            set_pgnum(sec._sectPr, "decimal", 1 if first_main else None)
            if first_main:
                page_header(sec)          # every later section inherits it
            first_main = False
        else:                                               # references, appendix, log
            set_pgnum(sec._sectPr, "lowerLetter", 1)
        width = TEXT_W_CM
        if role == "landscape":
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(PAGE_H_CM), Cm(PAGE_W_CM)
            sec.top_margin, sec.bottom_margin = Cm(LAND_MARGIN["top"]), Cm(LAND_MARGIN["bottom"])
            sec.left_margin, sec.right_margin = Cm(LAND_MARGIN["left"]), Cm(LAND_MARGIN["right"])
            sec.header_distance = sec.footer_distance = Cm(LAND_HF_CM)
            width = LAND_TEXT_W_CM
        # a footer is inherited from the previous section unless redefined; only the
        # first one and those whose page width changes need their own
        if width != previous_width:
            page_footer(sec, width)
            previous_width = width


# =============================================================================
# 7. Word: fit tables, update every field, export PDF
# =============================================================================
def build_equations(d) -> int:
    """Turn the LaTeX source in every equation table into a native OMML equation, using
    Word's own converter — the same one behind Insert > Equation with LaTeX input."""
    built = 0
    for t in d.Tables:
        if t.Title != "EQ":
            continue
        # Word fuses tables that touch, so consecutive equations arrive as one table
        # with a row per equation
        for row in range(1, t.Rows.Count + 1):
            cell = t.Cell(row, 2)
            for k in range(1, cell.Range.Paragraphs.Count + 1):
                r = cell.Range.Paragraphs(k).Range
                while r.Text and r.Text[-1] in "\r\x07":     # paragraph / end-of-cell mark
                    r.MoveEnd(1, -1)
                source = r.Text
                m = d.OMaths.Add(r).OMaths(1)
                m.BuildUp()
                if "\\" in m.Range.Text:
                    sys.exit("Word không dựng được công thức từ LaTeX:\n  " + source + "\n"
                             "Hãy đặt Word ở chế độ nhập LaTeX (Insert > Equation > "
                             "Conversions > LaTeX) rồi chạy lại.")
                built += 1
    return built


def word_finish(docx: Path, pdf: Path, eq_lines: int) -> None:
    import win32com.client as win32
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(docx))
        built = build_equations(d)
        if built != eq_lines:
            sys.exit(f"Word chỉ dựng {built}/{eq_lines} dòng công thức")
        print(f"   Word dựng {built} dòng công thức từ LaTeX")
        d.Fields.Update()
        d.Repaginate()
        for _ in range(2):                 # the TOC's own length moves pages once
            for toc in d.TablesOfContents:
                toc.Update()
            for tof in d.TablesOfFigures:
                tof.Update()
            d.Repaginate()
        # Only the stories that already exist: asking a section for a header it does not
        # have makes Word create an empty one, and a section that owns a header gets its
        # page border drawn around that header instead of around the text — which would
        # push the frame on the cover pages out towards the paper edge.
        for story in (6, 7, 8, 9, 10, 11):        # even/primary/first header and footer
            try:
                rng = d.StoryRanges(story)
            except Exception:
                continue
            while rng is not None:
                rng.Fields.Update()
                rng = rng.NextStoryRange
        d.Save()
        d.ExportAsFixedFormat(str(pdf), 17)
        pages = d.ComputeStatistics(2)
        print(f"   Word: {pages} trang, {d.TablesOfContents.Count} mục lục, "
              f"{d.TablesOfFigures.Count} danh mục hình/bảng")
        d.Close(False)
    finally:
        word.Quit()


# =============================================================================
def main() -> None:
    work = Path(tempfile.mkdtemp(prefix="kltn_docx_"))
    print("1. tiền xử lý Markdown + pandoc")
    raw = run_pandoc(preprocess(SRC_MD.read_text(encoding="utf-8"), work), work)

    print("2. định dạng theo mẫu của Khoa")
    doc = Document(str(raw))
    apply_styles(doc)
    n_tables = format_tables(doc)
    list_markers(doc)
    counts = body_pass(doc, heading_numbering(doc))
    cover_ends = insert_covers(doc)
    build_sections(doc, cover_ends, counts.pop("section_ends"))
    cp = doc.core_properties
    cp.title = "Phát hiện ảo giác cho mô hình ngôn ngữ lớn tăng cường truy xuất tiếng Việt"
    cp.author = "Nguyễn Ngọc Lân, Nguyễn Tấn Minh"
    cp.subject = "Khóa luận tốt nghiệp ngành Khoa học dữ liệu — IUH 2026"
    cp.language = "vi-VN"
    doc.save(str(OUT_DOCX))
    print(f"   {n_tables} bảng dữ liệu, {counts}")

    print("3. Word cập nhật mục lục, danh mục, số trang và xuất PDF")
    pdf = OUT_DOCX.with_suffix(".pdf")
    word_finish(OUT_DOCX, pdf, counts["dòng công thức"])
    print(f"   {OUT_DOCX.name}: {OUT_DOCX.stat().st_size / 1e6:.1f} MB, "
          f"{pdf.name}: {pdf.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
